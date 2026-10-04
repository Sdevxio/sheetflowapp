import posixpath
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

import numpy as np
import pandas as pd
import xlrd
from openpyxl import load_workbook

from app.services.errors import WorkbookReadError

_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_PKG = "{http://schemas.openxmlformats.org/package/2006/relationships}"
_OFFICE = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


@dataclass
class FormulaRef:
    formula: str | None
    cached: bool
    error: bool = False


@dataclass
class SheetPayload:
    name: str
    matrix: list[list[Any]]
    formulas: dict[tuple[int, int], FormulaRef] = field(default_factory=dict)
    merged_ranges: list[str] = field(default_factory=list)
    exceeds_row_limit: bool = False
    exceeds_column_limit: bool = False
    physical_rows_exact: bool = True
    column_count: int = 0
    merge_warning: str | None = None


def extract_workbook(path: Path, extension: str, *, max_rows: int, max_columns: int, max_sheets: int) -> list[SheetPayload]:
    try:
        if extension == ".xls":
            return _read_xls(path, max_rows=max_rows, max_columns=max_columns, max_sheets=max_sheets)
        if extension == ".xlsx":
            return _read_xlsx(path, max_rows=max_rows, max_columns=max_columns, max_sheets=max_sheets)
    except WorkbookReadError:
        raise
    except Exception as exc:
        raise WorkbookReadError("The workbook could not be read. It may be corrupted, encrypted, or incomplete.") from exc
    raise WorkbookReadError("Only .xls and .xlsx files can be imported.")


def _read_xlsx(path: Path, *, max_rows: int, max_columns: int, max_sheets: int) -> list[SheetPayload]:
    merges = _xlsx_merges(path)
    formulas = _xlsx_formulas(path, max_rows=max_rows, max_columns=max_columns)
    excel = pd.ExcelFile(path, engine="openpyxl")
    if len(excel.sheet_names) > max_sheets:
        raise WorkbookReadError(
            f"This workbook has {len(excel.sheet_names)} sheets, above the limit of {max_sheets}. No sheets were dropped."
        )
    sheets: list[SheetPayload] = []
    for name in excel.sheet_names:
        frame = pd.read_excel(
            excel,
            sheet_name=name,
            header=None,
            dtype=object,
            keep_default_na=False,
            na_filter=False,
            nrows=max_rows + 1,
        )
        sheets.append(_frame_to_sheet(name, frame, formulas.get(name, {}), merges.get(name, []), max_rows, max_columns))
    return sheets


def _read_xls(path: Path, *, max_rows: int, max_columns: int, max_sheets: int) -> list[SheetPayload]:
    formula_map, merge_map, merge_warning = _xls_metadata(path)
    excel = pd.ExcelFile(path, engine="xlrd")
    if len(excel.sheet_names) > max_sheets:
        raise WorkbookReadError(
            f"This workbook has {len(excel.sheet_names)} sheets, above the limit of {max_sheets}. No sheets were dropped."
        )
    sheets: list[SheetPayload] = []
    for name in excel.sheet_names:
        frame = pd.read_excel(
            excel,
            sheet_name=name,
            header=None,
            dtype=object,
            keep_default_na=False,
            na_filter=False,
            nrows=max_rows + 1,
        )
        sheet = _frame_to_sheet(name, frame, formula_map.get(name, {}), merge_map.get(name, []), max_rows, max_columns)
        if merge_warning and sheet.merge_warning is None:
            sheet.merge_warning = merge_warning
        sheets.append(sheet)
    return sheets


def _frame_to_sheet(
    name: str,
    frame: pd.DataFrame,
    formulas: dict[tuple[int, int], FormulaRef],
    merges: list[str],
    max_rows: int,
    max_columns: int,
) -> SheetPayload:
    row_count = int(frame.shape[0])
    column_count = int(frame.shape[1]) if row_count or frame.shape[1] else 0
    exceeds_rows = row_count > max_rows
    exceeds_cols = column_count > max_columns
    matrix = _matrix_from_frame(frame)
    resolved = _attach_cache_flags(matrix, formulas)
    return SheetPayload(
        name=name,
        matrix=matrix,
        formulas=resolved,
        merged_ranges=merges,
        exceeds_row_limit=exceeds_rows,
        exceeds_column_limit=exceeds_cols,
        physical_rows_exact=not exceeds_rows,
        column_count=column_count,
    )


def _matrix_from_frame(frame: pd.DataFrame) -> list[list[Any]]:
    matrix: list[list[Any]] = []
    if frame.empty and frame.shape[1] == 0:
        return matrix
    for record in frame.itertuples(index=False, name=None):
        matrix.append([_python_cell(value) for value in record])
    return matrix


def _python_cell(value: Any) -> Any:
    if value is None or value is pd.NA:
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, np.bool_):
        return bool(value)
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        if np.isnan(value):
            return None
        return float(value)
    if isinstance(value, float) and np.isnan(value):
        return None
    if isinstance(value, str):
        return None if value.strip() == "" else value
    if isinstance(value, (datetime, date, bool, int)):
        return value
    return value


def _attach_cache_flags(matrix: list[list[Any]], formulas: dict[tuple[int, int], FormulaRef]) -> dict[tuple[int, int], FormulaRef]:
    resolved: dict[tuple[int, int], FormulaRef] = {}
    for (row, col), ref in formulas.items():
        current = None
        if 1 <= row <= len(matrix) and 1 <= col <= len(matrix[row - 1]):
            current = matrix[row - 1][col - 1]
        if isinstance(current, str) and current.startswith("="):
            matrix[row - 1][col - 1] = None
            current = None
        cached = current is not None and ref.cached and not ref.error
        resolved[(row, col)] = FormulaRef(formula=ref.formula, cached=cached, error=ref.error)
    return resolved


def _xlsx_formulas(path: Path, *, max_rows: int, max_columns: int) -> dict[str, dict[tuple[int, int], FormulaRef]]:
    found: dict[str, dict[tuple[int, int], FormulaRef]] = {}
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        for sheet in workbook.worksheets:
            sheet_formulas: dict[tuple[int, int], FormulaRef] = {}
            for row in sheet.iter_rows(max_row=max_rows + 1, max_col=max_columns):
                for cell in row:
                    if cell.data_type == "f" and isinstance(cell.value, str):
                        sheet_formulas[(cell.row, cell.column)] = FormulaRef(formula=cell.value, cached=True)
            found[sheet.title] = sheet_formulas
    finally:
        workbook.close()
    return found


def _safe_zip_target(target: str) -> str:
    target = target.split("?", 1)[0].replace("\\", "/")
    if target.startswith("/"):
        normalized = posixpath.normpath(target.lstrip("/"))
    else:
        normalized = posixpath.normpath(posixpath.join("xl", target))
    if normalized.startswith("..") or normalized.startswith("/"):
        raise WorkbookReadError("The workbook contains an invalid internal path.")
    return normalized


def _xlsx_merges(path: Path) -> dict[str, list[str]]:
    merges: dict[str, list[str]] = {}
    with zipfile.ZipFile(path) as archive:
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        relationships = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {item.attrib.get("Id"): item.attrib.get("Target", "") for item in relationships}
        sheets = workbook.find(f"{_MAIN}sheets")
        if sheets is None:
            return merges
        for sheet in list(sheets):
            name = sheet.attrib.get("name", "")
            rel_id = sheet.attrib.get(f"{_OFFICE}id")
            target = targets.get(rel_id or "")
            if not target:
                merges[name] = []
                continue
            internal = _safe_zip_target(target)
            if internal not in archive.namelist():
                merges[name] = []
                continue
            root = ET.fromstring(archive.read(internal))
            refs = [node.attrib["ref"] for node in root.iter(f"{_MAIN}mergeCell") if node.attrib.get("ref")]
            merges[name] = refs
    return merges


def _xls_metadata(path: Path) -> tuple[dict[str, dict[tuple[int, int], FormulaRef]], dict[str, list[str]], str | None]:
    book = xlrd.open_workbook(path, formatting_info=True, on_demand=True)
    warning = None
    try:
        raw = bytes(book.mem)
        names = book.sheet_names()
        starts = list(book._sh_abs_posn)
        stream_end = min(len(raw), int(book.base) + int(book.stream_len))
        ends = starts[1:] + [stream_end]
        formulas: dict[str, dict[tuple[int, int], FormulaRef]] = {}
        merges: dict[str, list[str]] = {}
        for name, start, end in zip(names, starts, ends, strict=False):
            formulas[name] = _scan_biff_formulas(raw[start:end])
            sheet = book.sheet_by_name(name)
            ranges = []
            for rlo, rhi, clo, chi in getattr(sheet, "merged_cells", ()):
                ranges.append(f"{_col_letter(clo + 1)}{rlo + 1}:{_col_letter(chi)}{rhi}")
            merges[name] = ranges
        return formulas, merges, warning
    except Exception:
        warning = "Merged-cell ranges could not be read from this .xls file. Cell values were still read."
        return {}, {}, warning
    finally:
        book.release_resources()


def _scan_biff_formulas(blob: bytes) -> dict[tuple[int, int], FormulaRef]:
    found: dict[tuple[int, int], FormulaRef] = {}
    data = bytes(blob)
    index = 0
    while True:
        position = data.find(b"\x06\x00", index)
        if position < 0 or position + 4 > len(data):
            break
        length = int.from_bytes(data[position + 2 : position + 4], "little")
        index = position + 1
        if not (22 <= length <= 4096) or position + 4 + length > len(data):
            continue
        payload = data[position + 4 : position + 4 + length]
        row = int.from_bytes(payload[0:2], "little")
        col = int.from_bytes(payload[2:4], "little")
        if row > 65535 or col > 255 or payload[16:20] != b"\x00\x00\x00\x00":
            continue
        token_length = int.from_bytes(payload[20:22], "little")
        if token_length <= 0 or 22 + token_length != length:
            continue
        result = payload[6:14]
        error = False
        cached = True
        if result[6:8] == b"\xff\xff":
            kind = result[0]
            if kind == 3:
                cached = False
            elif kind == 2:
                cached = False
                error = True
            elif kind not in {0, 1}:
                continue
        found[(row + 1, col + 1)] = FormulaRef(formula=None, cached=cached, error=error)
    return found


def _col_letter(number: int) -> str:
    letters = ""
    while number:
        number, remainder = divmod(number - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters
