"""Order-details profile, mapping versions, and review records."""

from __future__ import annotations

import hashlib
from collections import defaultdict
from datetime import datetime, timezone
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from sqlalchemy import delete, select
from sqlalchemy.orm import Session
from sqlalchemy.orm.attributes import flag_modified

from app.models import (
    Import,
    ImportRow,
    MappingEntry,
    MappingVersion,
    UdsIssue,
    UdsLine,
    UdsLink,
    UdsObservation,
    UdsResolution,
)

ORDER_DETAILS_HEADERS = {
    "Report ID",
    "Order Name",
    "Facility",
    "Primary Insurance Name",
    "Patient Acct No",
    "Encounter Date",
    "Lab Attribute",
}

POSITIVE_VALUES = {"pos", "positive", "++positive++"}
INDETERMINATE_VALUES = {"inconsistent", "tnp", "unable to obtain"}
PRC_PRACTICE = "prc associates"
CATEGORY_BOTH = "screening and confirmation"
CATEGORY_SCREEN = "screening only"
CATEGORY_NONE = "not covered"


class UdsError(Exception):
    pass


def is_order_details(column_names: list[str]) -> bool:
    present = {str(name).strip() for name in column_names if str(name).strip()}
    return ORDER_DETAILS_HEADERS <= present


def line_role(attribute: str | None) -> str:
    text = (attribute or "").strip()
    if not text:
        return "order_line"
    lowered = text.lower()
    if "comment" in lowered or "notes" in lowered:
        return "comment"
    return "analyte"


def _key(value: str | None) -> str:
    return " ".join((value or "").split()).casefold()


def _text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _date(value) -> str | None:
    text = _text(value)
    if text is None:
        return None
    return text[:10]


def latest_mapping(session: Session) -> MappingVersion | None:
    return session.scalar(select(MappingVersion).order_by(MappingVersion.version_number.desc()).limit(1))


def load_mapping_workbook(payload: bytes, filename: str, session: Session) -> MappingVersion:
    digest = hashlib.sha256(payload).hexdigest()
    try:
        book = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise UdsError("The mapping workbook could not be read.") from exc
    frames = {}
    for name in book.sheetnames:
        frames[name.strip().casefold()] = pd.read_excel(BytesIO(payload), sheet_name=name, dtype=str)
    book.close()
    required = {
        "practices": "practices",
        "payer and categories": "payer and categories",
        "lab co and labs": "lab co and labs",
    }
    missing = [label for key, label in required.items() if key not in frames]
    if missing:
        raise UdsError("The mapping workbook needs Practices, PAYER AND CATEGORIES, and LAB CO and LABS.")
    practices = _practice_entries(frames["practices"])
    payers = _payer_entries(frames["payer and categories"])
    orders = _order_entries(frames["lab co and labs"])
    current = session.scalar(select(MappingVersion.version_number).order_by(MappingVersion.version_number.desc()).limit(1))
    version = MappingVersion(
        version_number=(current or 0) + 1,
        original_filename=filename,
        content_hash=digest,
        created_at=datetime.now(timezone.utc),
    )
    session.add(version)
    session.flush()
    for entry in practices + payers + orders:
        session.add(MappingEntry(version_id=version.id, **entry))
    session.flush()
    reclassify_all(session, version)
    return version


def _practice_entries(frame: pd.DataFrame) -> list[dict]:
    facility_col = _column(frame, "Facility")
    practice_col = _column(frame, "Practice")
    grouped: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for _, row in frame.iterrows():
        facility = _text(row.get(facility_col))
        practice = _text(row.get(practice_col))
        if facility is None:
            continue
        grouped[_key(facility)].append((facility, practice or ""))
    entries = []
    for pairs in grouped.values():
        practices = {practice for _, practice in pairs if practice}
        conflict = len(practices) > 1
        seen = set()
        for facility, practice in pairs:
            marker = practice.casefold()
            if marker in seen and not conflict:
                continue
            seen.add(marker)
            entries.append(
                {
                    "kind": "facility",
                    "source_key": facility,
                    "source_extra": None,
                    "mapped_value": practice or None,
                    "detail": None,
                    "conflict": conflict,
                }
            )
    return entries


def _payer_entries(frame: pd.DataFrame) -> list[dict]:
    payer_col = _column(frame, "Payer")
    category_col = _column(frame, "Payer UDS Category")
    id_col = _find_column(frame, "Insurance Payer ID")
    grouped: dict[str, list[tuple[str, str, str | None]]] = defaultdict(list)
    for _, row in frame.iterrows():
        payer = _text(row.get(payer_col))
        if payer is None:
            continue
        category = _text(row.get(category_col)) or ""
        payer_id = _text(row.get(id_col)) if id_col else None
        grouped[_key(payer)].append((payer, category, payer_id))
    entries = []
    for pairs in grouped.values():
        categories = {category for _, category, _ in pairs if category}
        conflict = len(categories) > 1
        seen = set()
        for payer, category, payer_id in pairs:
            marker = category.casefold()
            if marker in seen and not conflict:
                continue
            seen.add(marker)
            entries.append(
                {
                    "kind": "payer",
                    "source_key": payer,
                    "source_extra": payer_id,
                    "mapped_value": category or None,
                    "detail": None,
                    "conflict": conflict,
                }
            )
    return entries


def _order_entries(frame: pd.DataFrame) -> list[dict]:
    name_col = _column(frame, "Lab Order Name")
    company_col = _column(frame, "Lab Company")
    code_col = _find_column(frame, "Lab Code")
    type_col = _column(frame, "Lab Type")
    entries = []
    seen = set()
    for _, row in frame.iterrows():
        name = _text(row.get(name_col))
        if name is None or _key(name) in seen:
            continue
        seen.add(_key(name))
        lab_type = _text(row.get(type_col)) or ""
        order_class = "screening" if lab_type.casefold() == "presumptive" else "confirmation" if lab_type.casefold() == "confirmation" else "unclassified"
        entries.append(
            {
                "kind": "order",
                "source_key": name,
                "source_extra": _text(row.get(code_col)) if code_col else None,
                "mapped_value": _text(row.get(company_col)),
                "detail": {"lab_type": lab_type, "order_class": order_class},
                "conflict": False,
            }
        )
    return entries


def _column(frame: pd.DataFrame, name: str) -> str:
    found = _find_column(frame, name)
    if found is None:
        raise UdsError(f"The mapping workbook is missing the {name} column.")
    return found


def _find_column(frame: pd.DataFrame, name: str) -> str | None:
    wanted = name.casefold()
    for column in frame.columns:
        if str(column).strip().casefold() == wanted:
            return column
    return None


def mapping_lookup(session: Session, version: MappingVersion | None) -> dict:
    if version is None:
        return {"facility": {}, "payer": {}, "order": {}}
    rows = session.scalars(select(MappingEntry).where(MappingEntry.version_id == version.id)).all()
    lookup = {"facility": defaultdict(list), "payer": defaultdict(list), "order": defaultdict(list)}
    for row in rows:
        lookup[row.kind][_key(row.source_key)].append(row)
    return lookup


def build_observations(session: Session, job: Import) -> int:
    config = job.configuration or {}
    if config.get("profile") != "order_details":
        return 0
    version = latest_mapping(session)
    lookup = mapping_lookup(session, version)
    rows = session.scalars(
        select(ImportRow).where(ImportRow.import_id == job.id, ImportRow.status == "accepted").order_by(ImportRow.source_row)
    ).all()
    session.execute(delete(UdsIssue).where(UdsIssue.import_id == job.id))
    grouped: dict[str, list[ImportRow]] = defaultdict(list)
    for row in rows:
        data = row.transformed or {}
        report_id = _text(data.get("report_id"))
        if report_id is None:
            session.add(
                UdsIssue(
                    import_id=job.id,
                    observation_id=None,
                    report_id=None,
                    code="MISSING_REPORT_ID",
                    message="An accepted row has no Report ID, so it is not counted as an order.",
                    source_rows=[row.source_row],
                )
            )
            continue
        grouped[report_id].append(row)
    existing = session.scalars(select(UdsObservation).where(UdsObservation.import_id == job.id)).all()
    if existing:
        session.execute(delete(UdsLine).where(UdsLine.observation_id.in_([item.id for item in existing])))
        session.execute(delete(UdsObservation).where(UdsObservation.import_id == job.id))
        session.flush()
    created = []
    for report_id, members in grouped.items():
        observation = _observation_from_rows(job, report_id, members, version, lookup, session)
        created.append(observation)
    session.flush()
    choose_current(session, list(grouped))
    session.flush()
    refresh_links(session)
    return len(created)


def _observation_from_rows(job, report_id, members, version, lookup, session) -> UdsObservation:
    values = [row.transformed or {} for row in members]
    source_rows = [row.source_row for row in members]

    def stable(field: str) -> str | None:
        found = []
        for value in values:
            text = _text(value.get(field))
            if text is not None and text not in found:
                found.append(text)
        return found[0] if found else None

    def distinct(field: str) -> list[str]:
        found = []
        for value in values:
            text = _text(value.get(field))
            if text is not None and text not in found:
                found.append(text)
        return found

    payers = distinct("primary_insurance_name")
    icds = distinct("icd_code")
    observation = UdsObservation(
        import_id=job.id,
        report_id=report_id,
        order_name=stable("order_name"),
        facility=stable("facility"),
        provider=stable("appointment_servicing_provider"),
        payer_name=payers[0] if payers else None,
        account_id=stable("patient_acct_no"),
        encounter_date=_date(stable("encounter_date")),
        result_date=_date(stable("result_date")),
        reviewed_date=_date(stable("order_reviewed_date")),
        order_status=stable("order_status"),
        order_cancelled=stable("order_cancelled"),
        order_received=stable("order_received"),
        order_result=stable("order_result"),
        order_class="unclassified",
        lab_source="none",
        mapping_version_id=version.id if version else None,
        line_count=len(members),
        classification={},
    )
    session.add(observation)
    session.flush()
    analytes = comments = 0
    analyte_values = []
    for row, data in zip(members, values):
        attribute = _text(data.get("lab_attribute"))
        role = line_role(attribute)
        if role == "analyte":
            analytes += 1
            analyte_values.append(_text(data.get("lab_attribute_value")))
        elif role == "comment":
            comments += 1
        session.add(
            UdsLine(
                observation_id=observation.id,
                source_row=row.source_row,
                role=role,
                lab_attribute=attribute,
                lab_attribute_value=_text(data.get("lab_attribute_value")),
                icd_code=_text(data.get("icd_code")),
            )
        )
    observation.analyte_count = analytes
    observation.comment_count = comments
    _apply_mapping(observation, lookup, session, source_rows)
    if len(payers) > 1:
        _issue(session, job.id, observation, "INSURANCE_VARIANCE", "Primary insurance changes across lines of this order.", source_rows)
    if len(icds) > 1:
        _issue(session, job.id, observation, "ICD_VARIANCE", "ICD code changes across lines of this order.", source_rows)
    observation.classification = classify_observation(observation, analyte_values, session)
    return observation


def _apply_mapping(observation: UdsObservation, lookup: dict, session: Session, source_rows: list[int] | None = None) -> None:
    rows = source_rows or []
    facilities = lookup["facility"].get(_key(observation.facility), [])
    if observation.facility and not facilities:
        observation.practice = None
        _issue(session, observation.import_id, observation, "UNMAPPED_FACILITY", "This facility is not in the mapping workbook.", rows)
    elif facilities and any(item.conflict for item in facilities):
        names = sorted({item.mapped_value for item in facilities if item.mapped_value})
        observation.practice = None
        _issue(
            session,
            observation.import_id,
            observation,
            "FACILITY_CONFLICT",
            "This facility maps to more than one practice: " + ", ".join(names) + ".",
            rows,
        )
    elif facilities:
        observation.practice = facilities[0].mapped_value
    payers = lookup["payer"].get(_key(observation.payer_name), [])
    if observation.payer_name and not payers:
        observation.payer_category = "unknown"
        _issue(session, observation.import_id, observation, "UNMAPPED_PAYER", "This payer name is not in the mapping workbook. It is unknown, not 'not covered'.", rows)
    elif payers and any(item.conflict for item in payers):
        observation.payer_category = "unknown"
        _issue(session, observation.import_id, observation, "PAYER_CONFLICT", "This payer name maps to more than one category.", rows)
    elif payers:
        observation.payer_category = payers[0].mapped_value
    orders = lookup["order"].get(_key(observation.order_name), [])
    if not orders:
        observation.order_class = "unclassified"
        observation.lab_source = "none"
        _issue(session, observation.import_id, observation, "UNCLASSIFIED_ORDER", "This order name is not in the standard order list. It is unclassified, not non-UDS.", rows)
    else:
        detail = orders[0].detail or {}
        observation.order_class = detail.get("order_class") or "unclassified"
        observation.lab_from_order = orders[0].mapped_value
        observation.lab_code = orders[0].source_extra
        observation.lab_source = "order_name"
    _apply_resolution_overrides(observation, session)
    observation.expected_lab = _expected_lab(observation)


def _apply_resolution_overrides(observation: UdsObservation, session: Session) -> None:
    rows = session.scalars(select(UdsResolution).where(UdsResolution.active.is_(True))).all()
    for row in rows:
        if row.scope == "order" and row.report_id != observation.report_id:
            continue
        decision = row.decision or {}
        if row.code == "UNMAPPED_PAYER" and _key(decision.get("payer")) == _key(observation.payer_name):
            observation.payer_category = decision.get("category") or observation.payer_category
        if row.code == "UNMAPPED_FACILITY" and _key(decision.get("facility")) == _key(observation.facility):
            observation.practice = decision.get("practice") or observation.practice
        if row.code == "UNCLASSIFIED_ORDER" and _key(decision.get("order_name")) == _key(observation.order_name):
            observation.order_class = decision.get("order_class") or observation.order_class
            observation.lab_from_order = decision.get("lab") or observation.lab_from_order
            observation.lab_source = "review_resolution"


def _expected_lab(observation: UdsObservation) -> str | None:
    if _key(observation.practice) == PRC_PRACTICE:
        if _key(observation.payer_category) == CATEGORY_NONE:
            return "none"
        if _key(observation.payer_category) == CATEGORY_SCREEN and observation.order_class == "screening":
            return "Lighthouse"
        if _key(observation.payer_category) == CATEGORY_BOTH and observation.order_class == "confirmation":
            return "Lighthouse"
        if _key(observation.payer_category) == CATEGORY_BOTH and observation.order_class == "screening":
            return "none"
        return None
    if observation.practice:
        return "SYMPHONY"
    return None


def _issue(session, import_id, observation, code, message, source_rows) -> None:
    session.add(
        UdsIssue(
            import_id=import_id,
            observation_id=observation.id,
            report_id=observation.report_id,
            code=code,
            message=message,
            source_rows=source_rows,
        )
    )


def choose_current(session: Session, report_ids: list[str]) -> None:
    if not report_ids:
        return
    rows = session.scalars(select(UdsObservation).where(UdsObservation.report_id.in_(report_ids))).all()
    grouped: dict[str, list[UdsObservation]] = defaultdict(list)
    for row in rows:
        grouped[row.report_id].append(row)
    for report_id, observations in grouped.items():
        session.execute(delete(UdsIssue).where(UdsIssue.report_id == report_id, UdsIssue.code == "IDENTITY_CONFLICT"))
        for item in observations:
            item.is_current = False
            item.identity_conflict = False
        if len(observations) == 1:
            observations[0].is_current = True
            observations[0].selection_reason = "This is the only observation of this Report ID."
            continue
        ranked = sorted(observations, key=_precedence_key)
        best = ranked[-1]
        second = ranked[-2]
        if _precedence_key(best) > _precedence_key(second):
            best.is_current = True
            best.selection_reason = "This observation has a later result date or reviewed date than the other files."
            for item in observations:
                if item is not best:
                    item.selection_reason = "Another observation has a later result date or reviewed date."
            continue
        for item in observations:
            item.identity_conflict = True
            item.selection_reason = "The result and reviewed dates do not order these files. Both observations are kept."
            session.add(
                UdsIssue(
                    import_id=item.import_id,
                    observation_id=item.id,
                    report_id=report_id,
                    code="IDENTITY_CONFLICT",
                    message="Two observations of this Report ID cannot be ordered by result date or reviewed date.",
                    source_rows=[],
                )
            )


def _precedence_key(observation: UdsObservation) -> tuple[str, str]:
    return (observation.result_date or "", observation.reviewed_date or "")


def classify_observation(observation: UdsObservation, analyte_values: list[str | None], session: Session) -> dict:
    cancelled = _yes(observation.order_cancelled)
    lifecycle = "cancelled" if cancelled else "reviewed" if _key(observation.order_status) == "reviewed" else "open"
    positives = [value for value in analyte_values if _key(value) in POSITIVE_VALUES]
    indeterminate = [value for value in analyte_values if _key(value) in INDETERMINATE_VALUES]
    if positives:
        result = "positive"
        result_reason = "At least one analyte is on the explicit positive list. Other unknown analytes do not cancel that positive."
    elif indeterminate:
        result = "review_required"
        result_reason = "An analyte is indeterminate and no explicit positive is present. A negative result is not established."
    else:
        result = "not_established"
        result_reason = "No explicit positive is present. A negative result requires an agreed panel that is not defined yet."
    routing, destination, routing_reason, override_open = _routing(observation, result, cancelled)
    billing, billing_reason = _billing(observation, result, cancelled)
    confirmation = {
        "value": "not_applicable",
        "reason": "Confirmation follow-up applies to a positive screening.",
        "coverage": "",
    }
    if cancelled:
        confirmation = {"value": "not_applicable", "reason": "A cancelled screening is not followed for confirmation.", "coverage": ""}
    elif observation.order_class == "screening" and result == "positive":
        confirmation = {
            "value": "no_confirmation_in_available_data",
            "reason": "No confirmation found in available data.",
            "coverage": "",
        }
    compliance = _compliance(observation, result)
    return {
        "lifecycle": {"value": lifecycle, "reason": "Taken from Order Cancelled and Order Status."},
        "result": {"value": result, "reason": result_reason},
        "routing": {"value": routing, "destination": destination, "reason": routing_reason, "override_open": override_open},
        "billing": {"value": billing, "reason": billing_reason},
        "confirmation": confirmation,
        "compliance": compliance,
    }


def _routing(observation: UdsObservation, result: str, cancelled: bool) -> tuple[str, str | None, str, bool]:
    if cancelled:
        return "not_applicable", None, "A cancelled order is not routed.", False
    if not observation.practice:
        return "review_required", None, "No single practice is mapped, so no routing rule is applied.", False
    if _key(observation.practice) == PRC_PRACTICE:
        return "prc_matrix", observation.expected_lab, "The PRC matrix applies because the mapped practice is PRC Associates.", False
    return (
        "symphony_in_house",
        "SYMPHONY",
        "This mapped practice uses the SYMPHONY in-house rule. The payer category is recorded beside it. Whether that rule overrides the payer category is still open.",
        True,
    )


def _billing(observation: UdsObservation, result: str, cancelled: bool) -> tuple[str, str]:
    if cancelled:
        return "excluded_cancelled", "Cancelled orders stay visible and are excluded from billing forecasts."
    if not observation.practice or observation.payer_category in {None, "unknown"}:
        return "review_required", "Billing needs a practice and a known payer category."
    category = _key(observation.payer_category)
    if _key(observation.practice) == PRC_PRACTICE:
        if category == CATEGORY_NONE:
            return "not_eligible", "PRC matrix: this payer is not covered."
        if category == CATEGORY_SCREEN and result == "positive":
            return "not_eligible", "PRC matrix: a positive screening is not billed. The confirmation is the billable order when it exists."
        if category == CATEGORY_SCREEN and result != "positive":
            return "pending", "PRC matrix: a negative screening can be eligible, but this result is not established as negative."
        if category == CATEGORY_BOTH:
            return "pending", "PRC matrix: confirmation is the billable type. G-code readiness is not defined, so no billing forecast is produced."
        return "review_required", "The payer category is not one of the PRC matrix categories."
    if category in {CATEGORY_NONE, CATEGORY_SCREEN}:
        return "review_required", f"Payer category '{observation.payer_category}' conflicts with the SYMPHONY in-house rule. The override is not confirmed."
    if category == CATEGORY_BOTH:
        return "eligible", "Payer category allows screening and confirmation, and the practice rule is SYMPHONY in-house. No dollar amount is calculated."
    return "review_required", "The payer category is not recognized."


def _compliance(observation: UdsObservation, result: str) -> dict:
    if not observation.practice or observation.payer_category in {None, "unknown"}:
        return {"value": "insufficient_evidence", "reason": "Practice or payer mapping is missing, so compliance is not established."}
    if observation.order_class == "unclassified":
        return {"value": "confirmed_exception", "reason": "The order name is not in the standard order list."}
    if observation.order_class == "screening" and result == "positive":
        return {"value": "insufficient_evidence", "reason": "A positive screening has no accepted confirmation in the loaded data. That is not proof an order is missing elsewhere."}
    return {"value": "no_exception", "reason": "No confirmed exception is recorded for this order."}


def _yes(value: str | None) -> bool:
    return _key(value) in {"yes", "true", "y", "1"}


def refresh_links(session: Session) -> None:
    current = session.scalars(select(UdsObservation).where(UdsObservation.is_current.is_(True))).all()
    by_encounter: dict[tuple[str, str], list[UdsObservation]] = defaultdict(list)
    coverage = _coverage(session, current)
    for item in current:
        if item.account_id and item.encounter_date:
            by_encounter[(item.account_id, item.encounter_date)].append(item)
    preserved = {
        (link.screening_report_id, link.confirmation_report_id): link
        for link in session.scalars(select(UdsLink)).all()
        if link.status in {"accepted", "rejected"}
    }
    session.flush()
    session.execute(delete(UdsLink).where(UdsLink.status == "suggested"), execution_options={"synchronize_session": False})
    for encounter, orders in by_encounter.items():
        screenings = [
            item
            for item in orders
            if item.order_class == "screening"
            and not _yes(item.order_cancelled)
            and (item.classification or {}).get("result", {}).get("value") == "positive"
        ]
        confirmations = [item for item in orders if item.order_class == "confirmation" and not _yes(item.order_cancelled)]
        shared = len(screenings) > 1 and len(confirmations) >= 1
        for screening in screenings:
            classification = dict(screening.classification or {})
            confirmation = dict(classification.get("confirmation") or {})
            confirmation["coverage"] = coverage
            if shared:
                confirmation["value"] = "linkage_uncertain"
                confirmation["reason"] = "One confirmation is not credited to more than one screening unless a reviewer accepts that relationship."
            elif len(confirmations) == 1:
                candidate = confirmations[0]
                key = (screening.report_id, candidate.report_id)
                existing = preserved.get(key)
                if existing is None:
                    session.add(
                        UdsLink(
                            screening_report_id=screening.report_id,
                            confirmation_report_id=candidate.report_id,
                            status="suggested",
                            reason="One confirmation candidate shares this provisional encounter. The link is suggested and not accepted.",
                            created_at=datetime.now(timezone.utc),
                        )
                    )
                    confirmation["value"] = "suggested"
                    confirmation["reason"] = "One confirmation on the same provisional encounter is suggested. It is not accepted."
                elif existing.status == "accepted":
                    confirmation["value"] = "accepted"
                    confirmation["reason"] = "A reviewer accepted the same-encounter confirmation link."
                else:
                    confirmation["value"] = "rejected"
                    confirmation["reason"] = "A reviewer rejected the suggested confirmation link."
            elif len(confirmations) > 1:
                confirmation["value"] = "linkage_uncertain"
                confirmation["reason"] = "More than one confirmation shares this provisional encounter. No link was created."
            else:
                confirmation["value"] = "no_confirmation_in_available_data"
                confirmation["reason"] = f"No confirmation found in available data. Exports loaded: {coverage}."
            classification["confirmation"] = confirmation
            screening.classification = classification
            flag_modified(screening, "classification")
    for screening in current:
        if screening.order_class != "screening":
            continue
        if (screening.classification or {}).get("result", {}).get("value") != "positive":
            continue
        if not screening.account_id or not screening.encounter_date:
            classification = dict(screening.classification or {})
            confirmation = dict(classification.get("confirmation") or {})
            confirmation["value"] = "no_confirmation_in_available_data"
            confirmation["coverage"] = coverage
            confirmation["reason"] = f"No confirmation found in available data. A provisional encounter is missing, so no link was attempted. Exports loaded: {coverage}."
            classification["confirmation"] = confirmation
            screening.classification = classification
            flag_modified(screening, "classification")


def _coverage(session: Session, current: list[UdsObservation]) -> str:
    import_ids = {item.import_id for item in current}
    if not import_ids:
        return "none"
    jobs = session.scalars(select(Import).where(Import.id.in_(import_ids))).all()
    names = []
    for job in jobs:
        finished = job.finished_at.date().isoformat() if job.finished_at else "unknown date"
        names.append(f"{job.original_filename} ({finished})")
    return ", ".join(names) if names else "none"


def reclassify_all(session: Session, version: MappingVersion) -> None:
    lookup = mapping_lookup(session, version)
    observations = session.scalars(select(UdsObservation)).all()
    report_ids = []
    for observation in observations:
        session.execute(
            delete(UdsIssue).where(
                UdsIssue.observation_id == observation.id,
                UdsIssue.code.in_(["UNMAPPED_FACILITY", "FACILITY_CONFLICT", "UNMAPPED_PAYER", "PAYER_CONFLICT", "UNCLASSIFIED_ORDER"]),
            )
        )
        lines = session.scalars(select(UdsLine).where(UdsLine.observation_id == observation.id)).all()
        observation.mapping_version_id = version.id
        _apply_mapping(observation, lookup, session, [line.source_row for line in lines])
        lines = [line for line in lines if line.role == "analyte"]
        observation.classification = classify_observation(observation, [line.lab_attribute_value for line in lines], session)
        report_ids.append(observation.report_id)
    session.flush()
    choose_current(session, report_ids)
    refresh_links(session)


def public_observation(observation: UdsObservation) -> dict:
    return {
        "id": observation.id,
        "import_id": str(observation.import_id),
        "report_id": observation.report_id,
        "order_name": observation.order_name,
        "facility": observation.facility,
        "practice": observation.practice,
        "provider": observation.provider,
        "payer_name": observation.payer_name,
        "payer_category": observation.payer_category,
        "provisional_account": observation.account_id,
        "provisional_encounter_date": observation.encounter_date,
        "result_date": observation.result_date,
        "reviewed_date": observation.reviewed_date,
        "order_status": observation.order_status,
        "order_cancelled": observation.order_cancelled,
        "order_class": observation.order_class,
        "lab_from_order": observation.lab_from_order,
        "lab_code": observation.lab_code,
        "expected_lab": observation.expected_lab,
        "lab_source": observation.lab_source,
        "is_current": observation.is_current,
        "identity_conflict": observation.identity_conflict,
        "selection_reason": observation.selection_reason,
        "analyte_count": observation.analyte_count,
        "comment_count": observation.comment_count,
        "line_count": observation.line_count,
        "classification": observation.classification or {},
        "counting_unit": "order (one Report ID in one import)",
    }
