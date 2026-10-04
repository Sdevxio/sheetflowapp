"""Build the synthetic demonstration workbooks. These are not business records."""

from pathlib import Path

import xlwt
from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[2]
SAMPLES = ROOT / "samples"


def build_xlsx(path: Path) -> None:
    workbook = Workbook()
    orders = workbook.active
    orders.title = "Orders"
    orders.merge_cells("A1:F1")
    title = orders["A1"]
    title.value = "Synthetic demonstration workbook — not business data"
    title.font = Font(name="Arial", bold=True, size=14)
    orders["A2"] = ""
    headers = ["Account Code", "Order Date", "Amount", "Status", "Note", "Tax Rate"]
    for index, header in enumerate(headers, start=1):
        cell = orders.cell(3, index, header)
        cell.font = Font(name="Arial", bold=True)
    rows = [
        (4, "00123", "2024-01-15", 10.5, "open", "first", 0.08),
        (5, "00456", "01/02/2024", 20, "closed", None, "8%"),
        (6, "00123", "2024-03-01", None, "open", "missing amount", 0.05),
        (7, "00010", "not-a-date", "n/a", "pending", "bad values", None),
        (8, "", None, None, None, None, None),
        (9, "00010", "2024-04-01", 5, "open", "second account", 0.08),
        (10, "00123", "2024-01-15", 10.5, "open", "first", 0.08),
        (11, "Total", None, 999, None, None, None),
    ]
    for row_number, *values in rows:
        for index, value in enumerate(values, start=1):
            if value is None:
                continue
            orders.cell(row_number, index, value)

    comments = workbook.create_sheet("Comments")
    comments["A1"] = "Author"
    comments["B1"] = "Comment"
    comments["A1"].font = Font(name="Arial", bold=True)
    comments["B1"].font = Font(name="Arial", bold=True)
    comments["A2"] = "Ada Lovelace"
    comments["B2"] = "Synthetic note one"
    comments["A3"] = "Grace Hopper"
    comments["B3"] = "Synthetic note two"
    comments["A4"] = "Sam"
    injection = comments.cell(4, 2, "=1+1")
    injection.data_type = "s"
    orders_injection = orders.cell(9, 5, "=1+1")
    orders_injection.data_type = "s"

    calculations = workbook.create_sheet("Calculations")
    for index, header in enumerate(["Label", "Left", "Right", "Total"], start=1):
        cell = calculations.cell(1, index, header)
        cell.font = Font(name="Arial", bold=True)
    calculations["A2"] = "uncached"
    calculations["B2"] = 2
    calculations["C2"] = 3
    calculations["D2"] = "=B2+C2"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def build_xls(path: Path) -> None:
    book = xlwt.Workbook()
    sheet = book.add_sheet("Accounts")
    text = xlwt.easyxf(num_format_str="@")
    bold = xlwt.easyxf("font: bold on;")
    sheet.write(0, 0, "Account Code", bold)
    sheet.write(0, 1, "Label", bold)
    sheet.write(0, 2, "Computed", bold)
    sheet.write(1, 0, "00123", text)
    sheet.write(1, 1, "alpha")
    sheet.write(1, 2, xlwt.Formula("1+1"))
    sheet.write(2, 0, "00456", text)
    sheet.write(2, 1, "beta")
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(str(path))


def build_operations(path: Path) -> None:
    """Synthetic September workbook. Counts are produced by the import pipeline, not by the dashboard."""
    workbook = Workbook()
    orders = workbook.active
    orders.title = "Orders"
    for index, header in enumerate(["Order ID", "Date", "Region", "Status"], start=1):
        cell = orders.cell(1, index, header)
        cell.font = Font(name="Arial", bold=True)
    regions = ["North", "South", "East", "West"]
    row_number = 2

    def add(order_id: str | None, date: str, region: str, status: str) -> None:
        nonlocal row_number
        if order_id is not None:
            cell = orders.cell(row_number, 1, order_id)
            cell.number_format = "@"
        orders.cell(row_number, 2, date)
        orders.cell(row_number, 3, region)
        orders.cell(row_number, 4, status)
        row_number += 1

    sequence = 1
    for count, status in ((940, "Completed"), (200, "Pending"), (90, "Cancelled")):
        for _ in range(count):
            add(f"{sequence:06d}", "2024-09-15", regions[sequence % 4], status)
            sequence += 1
    for _ in range(12):
        add(f"{sequence:06d}", "not-a-date", "North", "Completed")
        sequence += 1
    for _ in range(6):
        add(None, "2024-09-16", "South", "Pending")
        sequence += 1

    regions_sheet = workbook.create_sheet("Regions")
    regions_sheet["A1"] = "Region"
    regions_sheet["B1"] = "Label"
    for index, name in enumerate(regions, start=2):
        regions_sheet.cell(index, 1, name)
        regions_sheet.cell(index, 2, f"{name} region")

    notes = workbook.create_sheet("Notes")
    notes["A1"] = "Note"
    notes["A2"] = "Synthetic demonstration workbook. Not business data."
    notes["A3"] = "September operations sample for SheetFlow."
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def main() -> None:
    build_xlsx(SAMPLES / "synthetic_demo.xlsx")
    build_xls(SAMPLES / "synthetic_demo.xls")
    build_operations(SAMPLES / "operations-september.xlsx")
    print(f"Wrote {SAMPLES / 'synthetic_demo.xlsx'}")
    print(f"Wrote {SAMPLES / 'synthetic_demo.xls'}")
    print(f"Wrote {SAMPLES / 'operations-september.xlsx'}")


if __name__ == "__main__":
    main()
